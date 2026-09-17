"""
core/excel_exporter.py
Schema-aware openpyxl workbook generator with:
  - Merged-range cell styling loop (borders on all cells in merged range)
  - Dynamic column map (no hardcoded column letters)
  - Native =SUM() formulas for Total Quantity Ordered and grand total rows
  - Two-tier colored headers, alternating row fills, auto-fit columns
"""
from __future__ import annotations

import io
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries

# ─────────────────────────────────────────────────────────────────────────────
# Palette helpers
# ─────────────────────────────────────────────────────────────────────────────

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color.lstrip("#"))


def _thin_border(color: str = "000000") -> Border:
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


PALETTE = {
    "master":    _fill("#86EFAC"),
    "ordered":   _fill("#67E8F9"),
    "fg":        _fill("#A7F3D0"),
    "pending":   _fill("#67E8F9"),
    "subheader": _fill("#F472B6"),
    "total_row": _fill("#FEF9C3"),
    "odd_row":   _fill("#FFFFFF"),
    "even_row":  _fill("#F8FAFC"),
}

TIER1_GROUPS = [
    {
        "label": "MASTER DATA",
        "cols":  ["Order Number", "Customer Name", "Order Date"],
        "fill":  "master",
    },
    {
        "label": "ORDERED QUANTITY",
        "cols":  [
            "TEMP (Ordered)", "LAMI (Ordered)", "IGU (Ordered)",
            "LAMI + IGU (Ordered)", "Total Quantity Ordered",
        ],
        "fill": "ordered",
    },
    {
        "label": "FINISHED GOODS",
        "cols":  ["Finished Goods Quantity"],
        "fill":  "fg",
    },
    {
        "label": "PENDING QUANTITY",
        "cols":  [
            "TEMP (Pending)", "LAMI (Pending)",
            "IGU (Pending)", "LAMI + IGU (Pending)",
        ],
        "fill": "pending",
    },
]

ORDERED_CATEGORY_COLS = [
    "TEMP (Ordered)", "LAMI (Ordered)",
    "IGU (Ordered)", "LAMI + IGU (Ordered)",
]

NUMERIC_COLS = {
    "TEMP (Ordered)", "LAMI (Ordered)", "IGU (Ordered)", "LAMI + IGU (Ordered)",
    "Total Quantity Ordered", "Finished Goods Quantity",
    "TEMP (Pending)", "LAMI (Pending)", "IGU (Pending)", "LAMI + IGU (Pending)",
}


# ─────────────────────────────────────────────────────────────────────────────
# Schema-aware column map
# ─────────────────────────────────────────────────────────────────────────────

def build_col_map(df: pd.DataFrame) -> dict[str, str]:
    """
    Map DataFrame column names -> Excel column letters (1-indexed, A-based).
    Works for any column order or subset — no hardcoded letters.
    """
    return {
        col_name: get_column_letter(col_idx + 1)
        for col_idx, col_name in enumerate(df.columns)
    }


# ─────────────────────────────────────────────────────────────────────────────
# Formula builders
# ─────────────────────────────────────────────────────────────────────────────

def _make_row_sum_formula(
    col_map: dict[str, str],
    source_cols: list[str],
    row_num: int,
) -> str:
    """
    Build a row-level SUM formula from source column names.
    - Contiguous columns  -> =SUM(D5:G5)
    - Non-contiguous      -> =D5+E5+F5
    - Single column       -> =D5
    - No matching columns -> literal 0
    """
    letters = [col_map[c] for c in source_cols if c in col_map]
    if not letters:
        return "0"
    if len(letters) == 1:
        return f"={letters[0]}{row_num}"

    indices = [column_index_from_string(ltr) for ltr in letters]
    is_contiguous = (indices == list(range(indices[0], indices[-1] + 1)))

    if is_contiguous:
        return f"=SUM({letters[0]}{row_num}:{letters[-1]}{row_num})"
    return "=" + "+".join(f"{ltr}{row_num}" for ltr in letters)


def _make_col_sum_formula(col_letter: str, start_row: int, end_row: int) -> str:
    """Column grand total formula: =SUM(D3:D52)"""
    return f"=SUM({col_letter}{start_row}:{col_letter}{end_row})"


# ─────────────────────────────────────────────────────────────────────────────
# Cell styling
# ─────────────────────────────────────────────────────────────────────────────

def _style_cell(
    cell,
    fill: Optional[PatternFill] = None,
    bold: bool = False,
    align: str = "center",
    number_format: str = "General",
) -> None:
    if fill:
        cell.fill = fill
    cell.font          = Font(bold=bold, name="Calibri", size=10)
    cell.alignment     = Alignment(horizontal=align, vertical="center", wrap_text=False)
    cell.border        = _thin_border()
    cell.number_format = number_format


def _apply_merged_range_style(
    ws, merge_ref: str, fill: PatternFill, bold: bool = True
) -> None:
    """
    Apply fill and borders to ALL cells in a merged range.
    """
    min_col, min_row, max_col, max_row = range_boundaries(merge_ref)
    border = _thin_border()
    from copy import copy

    for row_idx in range(min_row, max_row + 1):
        for col_idx in range(min_col, max_col + 1):
            cell           = ws.cell(row=row_idx, column=col_idx)
            cell.fill      = copy(fill)
            cell.border    = copy(border)
            cell.font      = Font(bold=bold, name="Calibri", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")


def _auto_fit_columns(
    ws, col_map: dict[str, str], df: pd.DataFrame
) -> None:
    """Set column widths based on maximum content length."""
    for col_name, col_letter in col_map.items():
        header_len = len(str(col_name))
        data_len = (
            df[col_name].astype(str).str.len().max()
            if col_name in df.columns and len(df) > 0 else 0
        )
        ws.column_dimensions[col_letter].width = min(
            max(header_len, int(data_len or 0)) + 4, 42
        )


# ─────────────────────────────────────────────────────────────────────────────
# Main exporter
# ─────────────────────────────────────────────────────────────────────────────

def export_to_excel(
    df: pd.DataFrame,
    sheet_title: str = "CDPL Report",
) -> bytes:
    """
    Generate a styled openpyxl workbook from a pivot DataFrame.

    Layout:
      Row 1   -> Tier 1 merged group headers (palette fills)
      Row 2   -> Tier 2 column subheaders (#F472B6)
      Row 3+  -> Data rows (alternating white / light-gray)
      Last    -> TOTAL summary row (#FEF9C3) with native =SUM() formulas

    Returns raw bytes for st.download_button().
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.freeze_panes = "A3"

    col_map    = build_col_map(df)
    DATA_START = 3
    data_end   = DATA_START + max(len(df) - 1, 0)
    total_row  = data_end + 1

    # ── Row 1: Tier 1 group headers (merged, with full-range styling loop) ──
    for group in TIER1_GROUPS:
        present_cols = [c for c in group["cols"] if c in col_map]
        if not present_cols:
            continue

        first_ltr = col_map[present_cols[0]]
        last_ltr  = col_map[present_cols[-1]]
        merge_ref = f"{first_ltr}1:{last_ltr}1"
        fill_obj  = PALETTE[group["fill"]]

        # Apply fill + border to EVERY cell in the merged range BEFORE merging
        _apply_merged_range_style(ws, merge_ref, fill=fill_obj, bold=True)
        ws.merge_cells(merge_ref)

        # Set value only on anchor (top-left) cell
        ws[f"{first_ltr}1"].value = group["label"]

    # ── Row 2: Tier 2 column subheaders ─────────────────────────────────────
    for col_name, col_letter in col_map.items():
        cell = ws[f"{col_letter}2"]
        cell.value = col_name
        _style_cell(cell, fill=PALETTE["subheader"], bold=True, align="center")

    # ── Rows 3..N: Data ─────────────────────────────────────────────────────
    for row_offset, (_, data_row) in enumerate(df.iterrows()):
        excel_row = DATA_START + row_offset
        row_fill  = PALETTE["odd_row"] if row_offset % 2 == 0 else PALETTE["even_row"]

        for col_name, col_letter in col_map.items():
            cell = ws[f"{col_letter}{excel_row}"]

            if col_name == "Total Quantity Ordered":
                cell.value = _make_row_sum_formula(
                    col_map, ORDERED_CATEGORY_COLS, excel_row
                )
                _style_cell(
                    cell, fill=row_fill, bold=True,
                    align="center", number_format="#,##0",
                )

            elif col_name == "Order Date":
                cell.value = data_row.get(col_name, "")
                _style_cell(cell, fill=row_fill, align="center")

            elif col_name in NUMERIC_COLS:
                raw = data_row.get(col_name, 0)
                cell.value = int(raw) if pd.notna(raw) else 0
                _style_cell(
                    cell, fill=row_fill, align="center", number_format="#,##0"
                )

            else:
                cell.value = (
                    str(data_row.get(col_name, ""))
                    if pd.notna(data_row.get(col_name, "")) else ""
                )
                _style_cell(cell, fill=row_fill, align="center")

    # ── TOTAL row ────────────────────────────────────────────────────────────
    for col_name, col_letter in col_map.items():
        cell = ws[f"{col_letter}{total_row}"]
        _style_cell(cell, fill=PALETTE["total_row"], bold=True)

        if col_name == "Order Number":
            cell.value     = "TOTAL"
            cell.alignment = Alignment(horizontal="center", vertical="center")

        elif col_name == "Total Quantity Ordered":
            ordered_letters = [
                col_map[c] for c in ORDERED_CATEGORY_COLS if c in col_map
            ]
            if ordered_letters:
                indices = [column_index_from_string(l) for l in ordered_letters]
                if indices == list(range(indices[0], indices[-1] + 1)):
                    cell.value = (
                        f"=SUM({ordered_letters[0]}{total_row}"
                        f":{ordered_letters[-1]}{total_row})"
                    )
                else:
                    cell.value = "=" + "+".join(
                        f"{l}{total_row}" for l in ordered_letters
                    )
            cell.number_format = "#,##0"
            cell.alignment     = Alignment(horizontal="center", vertical="center")

        elif col_name in NUMERIC_COLS:
            cell.value         = _make_col_sum_formula(col_letter, DATA_START, data_end)
            cell.number_format = "#,##0"
            cell.alignment     = Alignment(horizontal="center", vertical="center")

    # ── Polish ───────────────────────────────────────────────────────────────
    _auto_fit_columns(ws, col_map, df)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[total_row].height = 20

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
