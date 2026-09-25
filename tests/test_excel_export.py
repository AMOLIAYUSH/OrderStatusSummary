"""
tests/test_excel_export.py
openpyxl layout, palette fill, formula integrity, merged-range styling.
"""
from __future__ import annotations
import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from core.excel_exporter import (
    ORDERED_CATEGORY_COLS,
    build_col_map,
    export_to_excel,
    _make_row_sum_formula,
)


def _sample_df() -> pd.DataFrame:
    return pd.DataFrame([{
        "Order Number":          "RK000099001",
        "Customer Name":         "BAJRANG HOME SOLUTIONS",
        "Order Date":            "01-Sep-2026",
        "TEMP (Ordered)":        10,
        "LAMI (Ordered)":        5,
        "IGU (Ordered)":         3,
        "LAMI + IGU (Ordered)":  2,
        "Total Quantity Ordered": 20,
        "Finished Goods Quantity": 15,
        "TEMP (Pending)":        0,
        "LAMI (Pending)":        2,
        "IGU (Pending)":         0,
        "LAMI + IGU (Pending)":  2,
    }])


class TestColMap:
    def test_letters_are_correct(self):
        df = pd.DataFrame(columns=["Order Number", "Customer Name", "TEMP (Ordered)"])
        cm = build_col_map(df)
        assert cm["Order Number"]   == "A"
        assert cm["Customer Name"]  == "B"
        assert cm["TEMP (Ordered)"] == "C"


class TestRowSumFormula:
    def test_contiguous_four_cols_gives_sum_range(self):
        df = _sample_df()
        cm = build_col_map(df)
        formula = _make_row_sum_formula(cm, ORDERED_CATEGORY_COLS, row_num=5)
        assert formula.startswith("=SUM(") and ":" in formula

    def test_no_matching_cols_returns_zero(self):
        cm = {"Order Number": "A"}
        result = _make_row_sum_formula(cm, ORDERED_CATEGORY_COLS, row_num=5)
        assert result == "0"


class TestExcelExport:
    def test_valid_xlsx_output(self):
        data = export_to_excel(_sample_df())
        wb = load_workbook(io.BytesIO(data))
        assert "CTPL Report" in wb.sheetnames

    def test_tier1_master_data_fill(self):
        data = export_to_excel(_sample_df())
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        # A1 must be MASTER DATA -> #86EFAC
        fill_rgb = ws["A1"].fill.fgColor.rgb.upper()
        assert fill_rgb.endswith("86EFAC"), f"Got {fill_rgb}"

    def test_tier2_subheader_fill(self):
        data = export_to_excel(_sample_df())
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        # A2 must be Tier 2 subheader -> #F472B6
        fill_rgb = ws["A2"].fill.fgColor.rgb.upper()
        assert fill_rgb.endswith("F472B6"), f"Got {fill_rgb}"

    def test_total_qty_formula_in_data_row(self):
        data = export_to_excel(_sample_df())
        wb = load_workbook(io.BytesIO(data), data_only=False)
        ws = wb.active
        col_map = build_col_map(_sample_df())
        tot_col = col_map["Total Quantity Ordered"]
        cell_val = ws[f"{tot_col}3"].value
        assert isinstance(cell_val, str) and cell_val.startswith("=SUM("), \
            f"Expected =SUM(...) formula, got: {cell_val}"

    def test_total_row_has_col_sum_formula(self):
        data = export_to_excel(_sample_df())
        wb = load_workbook(io.BytesIO(data), data_only=False)
        ws = wb.active
        col_map = build_col_map(_sample_df())
        temp_col = col_map["TEMP (Ordered)"]
        total_row = 4   # 1 data row -> total at row 4
        cell_val = ws[f"{temp_col}{total_row}"].value
        assert isinstance(cell_val, str) and cell_val.startswith("=SUM("), \
            f"Expected =SUM(...) formula, got: {cell_val}"

    def test_merged_range_anchor_cell_styled(self):
        """The anchor cell in the Tier 1 MASTER DATA merged range must have a fill."""
        data = export_to_excel(_sample_df())
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        col_map = build_col_map(_sample_df())
        
        # Test only the anchor cell for the merged range
        ltr = col_map["Order Number"]
        cell = ws[f"{ltr}1"]
        fill_rgb = cell.fill.fgColor.rgb.upper()
        assert fill_rgb.endswith("86EFAC"), \
            f"Anchor cell {ltr}1 in MASTER DATA range missing fill, got {fill_rgb}"
